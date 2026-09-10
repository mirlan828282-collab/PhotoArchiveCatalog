import os
import re
import sys
import csv
import json
import math
import sqlite3
import shutil
import hashlib
import zipfile
from pathlib import Path
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

try:
    from PIL import Image, ImageTk, ImageDraw, ImageFont
except Exception:
    Image = None
    ImageTk = None
    ImageDraw = None
    ImageFont = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from docx import Document
except Exception:
    Document = None

from face_engine import FaceEngine

APP_NAME = 'Фотокаталог — Архив Президента Кыргызской Республики'
DEFAULT_PASSWORD = '12345'
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff', '.webp'}


def app_data_dir():
    if getattr(sys, 'frozen', False):
        base = os.environ.get('LOCALAPPDATA') or os.path.expanduser('~')
        root = Path(base) / 'PhotoArchiveCatalog'
    else:
        root = Path(__file__).resolve().parent / 'data'
    root.mkdir(parents=True, exist_ok=True)
    (root / 'photos').mkdir(parents=True, exist_ok=True)
    (root / 'thumbs').mkdir(parents=True, exist_ok=True)
    return root


DATA_DIR = app_data_dir()
DB_PATH = DATA_DIR / 'catalog.db'
PHOTOS_DIR = DATA_DIR / 'photos'
THUMBS_DIR = DATA_DIR / 'thumbs'


def safe_text(v):
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return v.strftime('%d.%m.%Y')
    return str(v).strip()


def normalize_key(s):
    return re.sub(r'[^0-9A-Za-zА-Яа-яЁё]+', '', safe_text(s)).casefold()


def whole_word_match(haystack, needle):
    haystack = safe_text(haystack)
    needle = safe_text(needle)
    if not needle:
        return True
    # Unicode-aware whole token/phrase boundaries. Hyphens in archive numbers remain searchable.
    pattern = r'(?<!\w)' + re.escape(needle) + r'(?!\w)'
    return re.search(pattern, haystack, flags=re.IGNORECASE | re.UNICODE) is not None


def copy_into_archive(src_path):
    src = Path(src_path)
    if not src.exists():
        raise FileNotFoundError(str(src))
    digest = hashlib.sha1((str(src.resolve()) + str(src.stat().st_mtime_ns)).encode('utf-8', 'ignore')).hexdigest()[:12]
    suffix = src.suffix.lower() or '.jpg'
    dest = PHOTOS_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{digest}{suffix}"
    shutil.copy2(str(src), str(dest))
    return str(dest)


def make_thumb(path, key, size=(86, 64)):
    if Image is None or not path or not os.path.exists(path):
        return None
    target = THUMBS_DIR / f'{key}_{size[0]}x{size[1]}.jpg'
    try:
        if not target.exists() or target.stat().st_mtime < Path(path).stat().st_mtime:
            img = Image.open(path)
            img.thumbnail(size)
            bg = Image.new('RGB', size, 'white')
            if img.mode != 'RGB':
                img = img.convert('RGB')
            bg.paste(img, ((size[0]-img.width)//2, (size[1]-img.height)//2))
            bg.save(str(target), 'JPEG', quality=88)
        return str(target)
    except Exception:
        return None


class CatalogDB:
    FIELDS = ('archive_no', 'description', 'shot_date', 'location', 'author', 'source', 'file_path')

    def __init__(self, path):
        self.path = str(path)
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute('PRAGMA foreign_keys=ON')
        self.create_schema()

    def create_schema(self):
        self.conn.executescript('''
        CREATE TABLE IF NOT EXISTS photos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            archive_no TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            shot_date TEXT NOT NULL DEFAULT '',
            location TEXT NOT NULL DEFAULT '',
            author TEXT NOT NULL DEFAULT '',
            source TEXT NOT NULL DEFAULT '',
            file_path TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_photos_archive_no ON photos(archive_no);
        CREATE INDEX IF NOT EXISTS idx_photos_shot_date ON photos(shot_date);
        CREATE INDEX IF NOT EXISTS idx_photos_author ON photos(author);
        ''')
        # v3 face index: several faces can belong to one photo.
        cols = [r['name'] for r in self.conn.execute("PRAGMA table_info(face_index)").fetchall()]
        if cols and 'id' not in cols:
            self.conn.execute('DROP TABLE IF EXISTS face_index')
        self.conn.executescript('''
        CREATE TABLE IF NOT EXISTS face_index (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            photo_id INTEGER NOT NULL,
            face_no INTEGER NOT NULL DEFAULT 0,
            embedding TEXT NOT NULL,
            indexed_at TEXT NOT NULL,
            FOREIGN KEY(photo_id) REFERENCES photos(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_face_photo ON face_index(photo_id);
        ''')
        self.conn.commit()

    def add(self, rec):
        now = datetime.now().isoformat(timespec='seconds')
        cur = self.conn.execute('''INSERT INTO photos
        (archive_no,description,shot_date,location,author,source,file_path,created_at,updated_at)
        VALUES (?,?,?,?,?,?,?,?,?)''',
        (rec.get('archive_no',''), rec.get('description',''), rec.get('shot_date',''), rec.get('location',''),
         rec.get('author',''), rec.get('source',''), rec.get('file_path',''), now, now))
        self.conn.commit(); return cur.lastrowid

    def update(self, photo_id, rec):
        now = datetime.now().isoformat(timespec='seconds')
        self.conn.execute('''UPDATE photos SET archive_no=?,description=?,shot_date=?,location=?,author=?,source=?,file_path=?,updated_at=? WHERE id=?''',
                          (rec.get('archive_no',''),rec.get('description',''),rec.get('shot_date',''),rec.get('location',''),
                           rec.get('author',''),rec.get('source',''),rec.get('file_path',''),now,photo_id))
        self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        self.conn.commit()

    def delete(self, photo_id):
        row = self.get(photo_id)
        self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        self.conn.execute('DELETE FROM photos WHERE id=?',(photo_id,))
        self.conn.commit(); return row

    def get(self, photo_id):
        return self.conn.execute('SELECT * FROM photos WHERE id=?',(photo_id,)).fetchone()

    def all(self):
        return self.conn.execute('SELECT * FROM photos ORDER BY id DESC').fetchall()

    def by_archive_no(self, archive_no):
        return self.conn.execute('SELECT * FROM photos WHERE lower(archive_no)=lower(?) ORDER BY id',(archive_no,)).fetchall()

    def search(self, text='', field='Все поля', whole_word=True):
        text = safe_text(text)
        rows = self.all()
        if not text:
            return rows
        fmap = {
            'Архивный номер': ['archive_no'], 'Описание': ['description'], 'Дата съёмки': ['shot_date'],
            'Место съёмки': ['location'], 'Автор': ['author'], 'Источник': ['source'],
            'Все поля': ['archive_no','description','shot_date','location','author','source']
        }
        fields = fmap.get(field, fmap['Все поля'])
        if whole_word:
            terms = [t for t in re.split(r'\s+', text) if t]
            result = []
            for r in rows:
                combined = ' '.join(safe_text(r[f]) for f in fields)
                if all(whole_word_match(combined, t) for t in terms):
                    result.append(r)
            return result
        low = text.casefold()
        return [r for r in rows if any(low in safe_text(r[f]).casefold() for f in fields)]

    def replace_faces(self, photo_id, embeddings):
        now = datetime.now().isoformat(timespec='seconds')
        self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        for i, emb in enumerate(embeddings):
            payload = json.dumps([float(x) for x in emb])
            self.conn.execute('INSERT INTO face_index(photo_id,face_no,embedding,indexed_at) VALUES (?,?,?,?)',
                              (photo_id,i,payload,now))
        self.conn.commit()

    def embeddings(self):
        return self.conn.execute('''SELECT f.id face_id,f.photo_id,f.face_no,f.embedding,p.*
                                    FROM face_index f JOIN photos p ON p.id=f.photo_id''').fetchall()

    def close(self): self.conn.close()


class PasswordDialog(simpledialog.Dialog):
    def body(self, master):
        ttk.Label(master, text='Введите пароль для перехода в защищённый раздел:').grid(row=0,column=0,padx=8,pady=(8,5))
        self.var = tk.StringVar()
        ent = ttk.Entry(master, textvariable=self.var, show='*', width=28)
        ent.grid(row=1,column=0,padx=8,pady=(0,8)); ent.focus_set(); return ent
    def apply(self): self.result = self.var.get()


class PhotoPopup(tk.Toplevel):
    def __init__(self, master, path, title='Фотография', max_size=(1100,760)):
        super().__init__(master); self.title(title); self.transient(master)
        self.ref = None
        if not path or not os.path.exists(path) or Image is None:
            ttk.Label(self,text='Файл фотографии не найден.',padding=30).pack(); return
        try:
            img=Image.open(path); img.thumbnail(max_size)
            self.ref=ImageTk.PhotoImage(img)
            ttk.Label(self,image=self.ref).pack(padx=8,pady=8)
        except Exception as exc:
            ttk.Label(self,text=f'Не удалось открыть фото:\n{exc}',padding=30).pack()


class Splash(tk.Toplevel):
    def __init__(self, master):
        super().__init__(master); self.overrideredirect(True); self.configure(bg='white')
        w,h=760,500
        sw,sh=self.winfo_screenwidth(),self.winfo_screenheight()
        self.geometry(f'{w}x{h}+{(sw-w)//2}+{(sh-h)//2}')
        canvas=tk.Canvas(self,width=w,height=h,bg='#eef6fb',highlightthickness=1,highlightbackground='#1a5578'); canvas.pack(fill='both',expand=True)
        # Stylised header: flag + emblem placeholder, so app works without external assets.
        canvas.create_rectangle(0,0,w,155,fill='#d81920',outline='')
        canvas.create_oval(300,20,460,180,fill='#16649a',outline='#f2cf3a',width=6)
        canvas.create_text(380,82,text='КЫРГЫЗ\nРЕСПУБЛИКАСЫ',fill='white',font=('Arial',16,'bold'),justify='center')
        canvas.create_oval(95,36,180,121,outline='#f2cf3a',width=7)
        for i in range(20):
            a=2*math.pi*i/20; x1=137+48*math.cos(a);y1=78+48*math.sin(a);x2=137+67*math.cos(a);y2=78+67*math.sin(a)
            canvas.create_line(x1,y1,x2,y2,fill='#f2cf3a',width=3)
        canvas.create_text(w//2,250,text='АРХИВ ПРЕЗИДЕНТА\nКЫРГЫЗСКОЙ РЕСПУБЛИКИ',fill='#123a58',font=('Arial',25,'bold'),justify='center')
        canvas.create_text(w//2,355,text='ФОТОКАТАЛОГ',fill='#16649a',font=('Arial',18,'bold'))
        self.pb=ttk.Progressbar(self,mode='indeterminate',length=430); self.pb.place(x=165,y=410); self.pb.start(12)
        canvas.create_text(w//2,455,text='Загрузка программы…',fill='#374957',font=('Arial',11))


class PhotoCatalogApp(tk.Tk):
    def __init__(self):
        super().__init__(); self.withdraw(); self.title(APP_NAME); self.geometry('1420x860'); self.minsize(1080,680)
        self.db=CatalogDB(DB_PATH); self.face=FaceEngine(); self.current_screen='viewer'; self.session_unlocked=False
        self.hover_popup=None; self.hover_after=None; self.current_rows=[]
        self.protocol('WM_DELETE_WINDOW',self.on_close)
        self.style=ttk.Style(self); self.style.configure('Treeview',rowheight=72)
        self.container=ttk.Frame(self); self.container.pack(fill='both',expand=True)
        self.screens={}
        for name, cls in [('viewer',ViewerScreen),('editor',EditorScreen),('catalog',CatalogScreen)]:
            f=cls(self.container,self); self.screens[name]=f; f.grid(row=0,column=0,sticky='nsew')
        self.container.rowconfigure(0,weight=1); self.container.columnconfigure(0,weight=1)
        splash=Splash(self)
        self.after(1400,lambda:self._finish_splash(splash))

    def _finish_splash(self,splash):
        try: splash.destroy()
        except Exception: pass
        self.deiconify(); self.show_screen('viewer', protected=False)

    def unlock(self):
        dlg=PasswordDialog(self,title='Ввод пароля')
        if dlg.result is None: return False
        if dlg.result != DEFAULT_PASSWORD:
            messagebox.showerror(APP_NAME,'Неверный пароль.'); return False
        self.session_unlocked=True; return True

    def show_screen(self,name,protected=None):
        if protected is None: protected = name in ('editor','catalog')
        if protected and not self.session_unlocked and not self.unlock(): return
        self.current_screen=name; self.screens[name].refresh(); self.screens[name].tkraise()

    def reset_lock(self): self.session_unlocked=False; self.show_screen('viewer',False)

    def popup_photo(self,path,title='Фотография'):
        PhotoPopup(self,path,title)

    def on_close(self):
        try:self.db.close()
        except Exception:pass
        self.destroy()


class BaseScreen(ttk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent); self.app=app; self.db=app.db; self.thumb_refs={}; self.rows=[]
        self.status_var=tk.StringVar(value='Готово')

    def nav(self, active):
        bar=ttk.Frame(self,padding=(10,8)); bar.pack(fill='x')
        ttk.Label(bar,text='АРХИВ ПРЕЗИДЕНТА КЫРГЫЗСКОЙ РЕСПУБЛИКИ',font=('Arial',12,'bold')).pack(side='left',padx=(0,20))
        ttk.Button(bar,text='1. Просмотр и поиск',command=lambda:self.app.show_screen('viewer',False)).pack(side='left',padx=3)
        ttk.Button(bar,text='2. Редактирование 🔒',command=lambda:self.app.show_screen('editor')).pack(side='left',padx=3)
        ttk.Button(bar,text='3. Каталог 🔒',command=lambda:self.app.show_screen('catalog')).pack(side='left',padx=3)
        ttk.Button(bar,text='Заблокировать',command=self.app.reset_lock).pack(side='right',padx=3)

    def status(self):
        ttk.Label(self,textvariable=self.status_var,relief='sunken',anchor='w',padding=(6,3)).pack(fill='x',side='bottom')

    def refresh(self): pass


class ViewerScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app); self.nav('viewer')
        body=ttk.Panedwindow(self,orient='horizontal'); body.pack(fill='both',expand=True,padx=10,pady=(0,8))
        left=ttk.Frame(body,padding=6); right=ttk.Frame(body,padding=8); body.add(left,weight=2); body.add(right,weight=3)
        search=ttk.LabelFrame(left,text='Поиск',padding=8); search.pack(fill='x')
        self.field=tk.StringVar(value='Все поля'); ttk.Combobox(search,textvariable=self.field,state='readonly',values=['Все поля','Архивный номер','Описание','Дата съёмки','Место съёмки','Автор','Источник']).pack(fill='x',pady=(0,6))
        self.q=tk.StringVar(); ent=ttk.Entry(search,textvariable=self.q); ent.pack(fill='x'); ent.bind('<Return>',lambda e:self.refresh())
        self.whole=tk.BooleanVar(value=True); ttk.Checkbutton(search,text='Искать только целое слово',variable=self.whole).pack(anchor='w',pady=6)
        bb=ttk.Frame(search);bb.pack(fill='x');ttk.Button(bb,text='Найти',command=self.refresh).pack(side='left');ttk.Button(bb,text='Сбросить',command=self.reset).pack(side='left',padx=5)
        ttk.Button(search,text='Поиск по лицу',command=self.search_face).pack(fill='x',pady=(8,0))

        listbox=ttk.LabelFrame(left,text='Результаты',padding=4); listbox.pack(fill='both',expand=True,pady=(8,0))
        self.tree=ttk.Treeview(listbox,columns=('no','date','desc'),show='headings',selectmode='browse')
        for c,t,w in [('no','Архивный №',130),('date','Дата',100),('desc','Описание',360)]:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='w')
        ys=ttk.Scrollbar(listbox,orient='vertical',command=self.tree.yview); xs=ttk.Scrollbar(listbox,orient='horizontal',command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set); self.tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew');listbox.rowconfigure(0,weight=1);listbox.columnconfigure(0,weight=1)
        self.tree.bind('<<TreeviewSelect>>',self.select)

        self.preview=ttk.Label(right,text='Выберите запись',anchor='center'); self.preview.pack(fill='both',expand=True)
        self.preview_ref=None; self.info=tk.Text(right,height=11,wrap='word',state='disabled'); self.info.pack(fill='x',pady=(8,0))
        ttk.Button(right,text='Открыть фотографию крупно',command=self.open_current).pack(anchor='e',pady=6)
        self.current=None; self.status()

    def reset(self): self.q.set('');self.refresh()
    def refresh(self):
        self.rows=self.db.search(self.q.get(),self.field.get(),self.whole.get());
        for i in self.tree.get_children(): self.tree.delete(i)
        for r in self.rows:
            desc=safe_text(r['description']).replace('\n',' '); self.tree.insert('', 'end',iid=str(r['id']),values=(r['archive_no'],r['shot_date'],desc[:120]))
        self.status_var.set(f'Найдено записей: {len(self.rows)}')
        if self.rows:
            self.tree.selection_set(str(self.rows[0]['id'])); self.select()

    def select(self,_e=None):
        s=self.tree.selection();
        if not s:return
        r=self.db.get(int(s[0])); self.current=r
        self.show_photo(r['file_path'])
        txt=(f"Архивный номер: {r['archive_no']}\nДата съёмки: {r['shot_date']}\nМесто съёмки: {r['location']}\n"
             f"Автор: {r['author']}\nИсточник поступления: {r['source']}\n\nОписание:\n{r['description']}")
        self.info.configure(state='normal'); self.info.delete('1.0','end');self.info.insert('1.0',txt);self.info.configure(state='disabled')

    def show_photo(self,path):
        self.preview_ref=None
        if not path or not os.path.exists(path) or Image is None:self.preview.configure(image='',text='Фотография отсутствует');return
        try:
            img=Image.open(path);img.thumbnail((760,500));self.preview_ref=ImageTk.PhotoImage(img);self.preview.configure(image=self.preview_ref,text='')
        except Exception:self.preview.configure(image='',text='Не удалось открыть фотографию')
    def open_current(self):
        if self.current:self.app.popup_photo(self.current['file_path'],self.current['archive_no'])

    def search_face(self):
        if not self.app.face.available: messagebox.showerror(APP_NAME,'Поиск по лицу недоступен:\n'+self.app.face.error);return
        path=filedialog.askopenfilename(title='Выберите фотографию лица',filetypes=[('Изображения','*.jpg *.jpeg *.png *.bmp *.tif *.tiff')]);
        if not path:return
        try:q=self.app.face.embedding(path)
        except Exception as exc:messagebox.showerror(APP_NAME,str(exc));return
        best={}
        for r in self.db.embeddings():
            try:
                score=self.app.face.cosine(q,json.loads(r['embedding'])); best[r['photo_id']]=max(best.get(r['photo_id'],-1),score)
            except Exception:pass
        ids=[pid for pid,sc in sorted(best.items(),key=lambda x:x[1],reverse=True) if sc>=0.36][:200]
        self.rows=[self.db.get(pid) for pid in ids]
        for i in self.tree.get_children():self.tree.delete(i)
        for r in self.rows:self.tree.insert('', 'end',iid=str(r['id']),values=(r['archive_no'],r['shot_date'],safe_text(r['description'])[:120]))
        self.status_var.set(f'Поиск по лицу: найдено {len(self.rows)}')
        if not ids:messagebox.showinfo(APP_NAME,'Совпадений нет. В окне редактирования сначала нажмите «Индексировать лица».')


class EditorScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app);self.nav('editor');self.current_id=None;self.current_path='';self.preview_ref=None
        top=ttk.Frame(self,padding=(10,0,10,6));top.pack(fill='x')
        ttk.Button(top,text='Новая карточка',command=self.new).pack(side='left')
        ttk.Button(top,text='Импорт Access',command=self.import_access).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт Excel/CSV',command=self.import_table).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт Word DOCX',command=self.import_word).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт фото по архивным №',command=self.import_photos_by_number).pack(side='left',padx=12)
        ttk.Button(top,text='Индексировать лица',command=self.index_faces).pack(side='left',padx=4)
        ttk.Button(top,text='Резервная копия',command=self.backup).pack(side='right')
        body=ttk.Panedwindow(self,orient='horizontal');body.pack(fill='both',expand=True,padx=10,pady=(0,8))
        lf=ttk.Frame(body,padding=5);rf=ttk.Frame(body,padding=8);body.add(lf,weight=2);body.add(rf,weight=3)
        self.q=tk.StringVar(); sr=ttk.Frame(lf);sr.pack(fill='x');ttk.Entry(sr,textvariable=self.q).pack(side='left',fill='x',expand=True);ttk.Button(sr,text='Найти',command=self.refresh).pack(side='left',padx=4)
        self.tree=ttk.Treeview(lf,columns=('no','date','desc'),show='headings',selectmode='browse')
        for c,t,w in [('no','Архивный №',130),('date','Дата',95),('desc','Описание',350)]:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='w')
        ys=ttk.Scrollbar(lf,orient='vertical',command=self.tree.yview);self.tree.configure(yscrollcommand=ys.set);self.tree.pack(side='left',fill='both',expand=True,pady=(6,0));ys.pack(side='right',fill='y',pady=(6,0));self.tree.bind('<<TreeviewSelect>>',self.select)
        self.entries={}; form=rf
        row=0
        for k,label in [('archive_no','Архивный номер'),('shot_date','Дата съёмки'),('location','Место съёмки'),('author','Автор съёмки'),('source','Источник поступления')]:
            ttk.Label(form,text=label).grid(row=row,column=0,sticky='w');e=ttk.Entry(form);e.grid(row=row+1,column=0,sticky='ew',pady=(0,6));self.entries[k]=e;row+=2
        ttk.Label(form,text='Описание фотографии').grid(row=row,column=0,sticky='w');row+=1;self.desc=tk.Text(form,height=6,wrap='word');self.desc.grid(row=row,column=0,sticky='nsew',pady=(0,6));row+=1
        pf=ttk.LabelFrame(form,text='Фотография',padding=6);pf.grid(row=row,column=0,sticky='nsew');self.preview=ttk.Label(pf,text='Фото не выбрано',anchor='center');self.preview.pack(fill='both',expand=True)
        pb=ttk.Frame(pf);pb.pack(fill='x',pady=5);ttk.Button(pb,text='Выбрать/заменить фото',command=self.choose_photo).pack(side='left');ttk.Button(pb,text='Удалить фото',command=self.remove_photo).pack(side='left',padx=5);ttk.Button(pb,text='Открыть',command=lambda:self.app.popup_photo(self.current_path)).pack(side='left');row+=1
        act=ttk.Frame(form);act.grid(row=row,column=0,sticky='ew',pady=6);ttk.Button(act,text='Сохранить',command=self.save).pack(side='left');ttk.Button(act,text='Удалить карточку',command=self.delete).pack(side='left',padx=6);ttk.Button(act,text='Отмена/очистить',command=self.new).pack(side='left')
        form.columnconfigure(0,weight=1);form.rowconfigure(row-1,weight=1);self.status()

    def refresh(self):
        self.rows=self.db.search(self.q.get(),'Все поля',True)
        for i in self.tree.get_children():self.tree.delete(i)
        for r in self.rows:self.tree.insert('', 'end',iid=str(r['id']),values=(r['archive_no'],r['shot_date'],safe_text(r['description'])[:100]))
        self.status_var.set(f'Записей: {len(self.rows)}')
    def select(self,_e=None):
        s=self.tree.selection();
        if not s:return
        r=self.db.get(int(s[0]));self.current_id=r['id'];self.current_path=r['file_path'] or ''
        for k,e in self.entries.items():e.delete(0,'end');e.insert(0,r[k] or '')
        self.desc.delete('1.0','end');self.desc.insert('1.0',r['description'] or '');self.show_preview()
    def new(self):
        self.current_id=None;self.current_path='';
        for e in self.entries.values():e.delete(0,'end')
        self.desc.delete('1.0','end');self.preview.configure(image='',text='Фото не выбрано');self.preview_ref=None
    def show_preview(self):
        self.preview_ref=None
        if not self.current_path or not os.path.exists(self.current_path) or Image is None:self.preview.configure(image='',text='Фото не выбрано');return
        try:
            img=Image.open(self.current_path);img.thumbnail((650,330));self.preview_ref=ImageTk.PhotoImage(img);self.preview.configure(image=self.preview_ref,text='')
        except Exception:self.preview.configure(image='',text='Ошибка изображения')
    def choose_photo(self):
        p=filedialog.askopenfilename(filetypes=[('Изображения','*.jpg *.jpeg *.png *.bmp *.tif *.tiff *.webp')]);
        if p:self.current_path=p;self.show_preview()
    def remove_photo(self):self.current_path='';self.show_preview()
    def collect(self):
        d={k:e.get().strip() for k,e in self.entries.items()};d['description']=self.desc.get('1.0','end').strip();d['file_path']=self.current_path;return d
    def save(self):
        d=self.collect();
        if not d['archive_no']:messagebox.showwarning(APP_NAME,'Укажите архивный номер.');return
        old=self.db.get(self.current_id) if self.current_id else None
        p=d['file_path']
        if p and os.path.exists(p):
            try:
                if old and old['file_path']==p:d['file_path']=p
                elif str(Path(p).resolve()).startswith(str(PHOTOS_DIR.resolve())):d['file_path']=p
                else:d['file_path']=copy_into_archive(p)
            except Exception as exc:messagebox.showerror(APP_NAME,str(exc));return
        if self.current_id:self.db.update(self.current_id,d)
        else:self.current_id=self.db.add(d)
        self.current_path=d['file_path'];self.refresh();self.status_var.set('Сохранено.')
    def delete(self):
        if not self.current_id:return
        if messagebox.askyesno(APP_NAME,'Удалить карточку из базы?'):
            self.db.delete(self.current_id);self.new();self.refresh()

    def _pick(self,row,aliases):
        norm={re.sub(r'[\s._№#:-]+',' ',safe_text(k).casefold()).strip():v for k,v in row.items()}
        for a in aliases:
            aa=re.sub(r'[\s._№#:-]+',' ',a.casefold()).strip()
            if aa in norm:return safe_text(norm[aa])
        return ''

    def _import_dict_rows(self,rows,base_dir):
        aliases={
        'archive_no':['архивный номер','архивный №','арх №','арх. №','архивный номер фото','номер','archive no','archive_no','шифр'],
        'description':['описание','описание фотографии','содержание','аннотация','description','сюжет'],
        'shot_date':['дата съемки','дата съёмки','дата','дата фото','shot date','shot_date'],
        'location':['место съемки','место съёмки','место','location'],
        'author':['автор съемки','автор съёмки','автор','фотограф','author'],
        'source':['источник поступления','источник','source','поступление'],
        'file_path':['фото','фотография','путь к фото','файл','file path','file_path','имя файла']}
        added=updated=skipped=0
        for row in rows:
            rec={k:self._pick(row,v) for k,v in aliases.items()}
            if not rec['archive_no']:
                skipped+=1;continue
            fp=rec.get('file_path','')
            if fp:
                p=Path(fp)
                if not p.is_absolute():p=Path(base_dir)/p
                if p.exists() and p.suffix.lower() in IMAGE_EXTS:rec['file_path']=copy_into_archive(p)
                else:rec['file_path']=''
            existing=self.db.by_archive_no(rec['archive_no'])
            # Upsert metadata into first matching record; never overwrite an existing photo with blank.
            if existing:
                old=existing[0]; merged={k:(rec[k] if rec.get(k) else old[k]) for k in CatalogDB.FIELDS}
                self.db.update(old['id'],merged);updated+=1
            else:self.db.add(rec);added+=1
        return added,updated,skipped

    def import_table(self):
        path=filedialog.askopenfilename(title='Импорт Excel / CSV',filetypes=[('Excel/CSV','*.xlsx *.xlsm *.csv'),('Все файлы','*.*')]);
        if not path:return
        try:
            rows=[]
            if Path(path).suffix.lower() in ('.xlsx','.xlsm'):
                if load_workbook is None:raise RuntimeError('Не установлен openpyxl.')
                ws=load_workbook(path,read_only=True,data_only=True).active;data=list(ws.iter_rows(values_only=True))
                if not data:raise ValueError('Таблица пустая.')
                headers=[safe_text(x) for x in data[0]];rows=[dict(zip(headers,v)) for v in data[1:] if any(x is not None for x in v)]
            else:
                try:f=open(path,'r',encoding='utf-8-sig',newline='')
                except UnicodeDecodeError:f=open(path,'r',encoding='cp1251',newline='')
                with f:rows=list(csv.DictReader(f))
            a,u,s=self._import_dict_rows(rows,Path(path).parent);self.refresh();messagebox.showinfo(APP_NAME,f'Готово.\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {s}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка импорта:\n{exc}')

    def import_word(self):
        path=filedialog.askopenfilename(title='Импорт Word',filetypes=[('Word DOCX','*.docx'),('Все файлы','*.*')]);
        if not path:return
        try:
            if Document is None:raise RuntimeError('Не установлен python-docx.')
            doc=Document(path);rows=[]
            for table in doc.tables:
                if not table.rows:continue
                headers=[safe_text(c.text) for c in table.rows[0].cells]
                for tr in table.rows[1:]:
                    vals=[safe_text(c.text) for c in tr.cells]
                    if any(vals):rows.append(dict(zip(headers,vals)))
            if not rows:raise ValueError('В документе Word не найдено таблиц с данными.')
            a,u,s=self._import_dict_rows(rows,Path(path).parent);self.refresh();messagebox.showinfo(APP_NAME,f'Word импортирован.\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {s}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка импорта Word:\n{exc}')

    def import_access(self):
        path=filedialog.askopenfilename(title='Импорт Microsoft Access',filetypes=[('Access','*.accdb *.mdb')]);
        if not path:return
        try:
            import pyodbc
            drivers=[d for d in pyodbc.drivers() if 'Access Driver' in d]
            if not drivers:raise RuntimeError('Не установлен Microsoft Access Database Engine (ODBC).')
            conn=pyodbc.connect(r'DRIVER={'+drivers[-1]+r'};DBQ='+path+';')
            tables=[r.table_name for r in conn.cursor().tables(tableType='TABLE') if not r.table_name.startswith('MSys')]
            if not tables:raise RuntimeError('Таблицы не найдены.')
            table=tables[0]
            if len(tables)>1:
                table=simpledialog.askstring(APP_NAME,'Таблицы: '+', '.join(tables[:15])+'\nВведите имя таблицы:',initialvalue=tables[0]) or ''
                if table not in tables:return
            cur=conn.cursor();cur.execute(f'SELECT * FROM [{table}]');cols=[d[0] for d in cur.description];rows=[dict(zip(cols,r)) for r in cur.fetchall()];conn.close()
            a,u,s=self._import_dict_rows(rows,Path(path).parent);self.refresh();messagebox.showinfo(APP_NAME,f'Таблица: {table}\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {s}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка Access:\n{exc}')

    def _match_archive_for_filename(self,stem,index):
        key=normalize_key(stem)
        if key in index:return index[key]
        # Common forms: ARCHIVENO_1, ARCHIVENO(2), ARCHIVENO-copy
        candidates=[]
        for k,rows in index.items():
            if not k:continue
            if key.startswith(k) or k in key:candidates.append((len(k),rows))
        if not candidates:return None
        candidates.sort(key=lambda x:x[0],reverse=True);return candidates[0][1]

    def import_photos_by_number(self):
        folder=filedialog.askdirectory(title='Выберите папку с фотографиями');
        if not folder:return
        recursive=messagebox.askyesno(APP_NAME,'Искать фотографии также во вложенных папках?')
        files=[]
        it=Path(folder).rglob('*') if recursive else Path(folder).glob('*')
        for p in it:
            if p.is_file() and p.suffix.lower() in IMAGE_EXTS:files.append(p)
        if not files:messagebox.showinfo(APP_NAME,'Фотографии не найдены.');return
        index={}
        for r in self.db.all():index.setdefault(normalize_key(r['archive_no']),[]).append(r)
        attached=cloned=unmatched=errors=0
        for n,p in enumerate(files,1):
            matches=self._match_archive_for_filename(p.stem,index)
            if not matches:unmatched+=1;continue
            try:
                # Prefer a matching card without a photo. Otherwise clone metadata for an additional photo.
                target=next((r for r in matches if not r['file_path']),None)
                dest=copy_into_archive(p)
                if target:
                    rec={k:target[k] for k in CatalogDB.FIELDS};rec['file_path']=dest;self.db.update(target['id'],rec);attached+=1
                else:
                    src=matches[0];rec={k:src[k] for k in CatalogDB.FIELDS};rec['file_path']=dest;self.db.add(rec);cloned+=1
            except Exception:errors+=1
            if n%25==0:self.status_var.set(f'Импорт фото: {n}/{len(files)}');self.update_idletasks()
        self.refresh();messagebox.showinfo(APP_NAME,f'Импорт фотографий завершён.\nПривязано к пустым карточкам: {attached}\nДополнительных карточек: {cloned}\nНе найден архивный номер: {unmatched}\nОшибок: {errors}')

    def index_faces(self):
        if not self.app.face.available:messagebox.showerror(APP_NAME,'Поиск по лицу недоступен:\n'+self.app.face.error);return
        rows=self.db.all();photos=faces=noface=errors=0
        for i,r in enumerate(rows,1):
            try:
                if r['file_path'] and os.path.exists(r['file_path']):
                    embs=self.app.face.embeddings(r['file_path']);self.db.replace_faces(r['id'],embs);photos+=1;faces+=len(embs)
                else:errors+=1
            except ValueError:self.db.replace_faces(r['id'],[]);noface+=1
            except Exception:errors+=1
            self.status_var.set(f'Индексация лиц: {i}/{len(rows)}');self.update_idletasks()
        messagebox.showinfo(APP_NAME,f'Индексация завершена.\nФотографий обработано: {photos}\nЛиц проиндексировано: {faces}\nБез лиц: {noface}\nОшибок: {errors}')
        self.status_var.set('Готово')

    def backup(self):
        target=filedialog.asksaveasfilename(defaultextension='.zip',initialfile=f"PhotoArchive_backup_{datetime.now():%Y%m%d_%H%M}.zip",filetypes=[('ZIP','*.zip')]);
        if not target:return
        try:
            self.db.conn.commit()
            with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
                for p in DATA_DIR.rglob('*'):
                    if p.is_file():z.write(p,p.relative_to(DATA_DIR))
            messagebox.showinfo(APP_NAME,'Резервная копия создана.')
        except Exception as exc:messagebox.showerror(APP_NAME,str(exc))


class CatalogScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app);self.nav('catalog');self.thumb_refs={};self.thumb_size=tk.IntVar(value=72)
        top=ttk.Frame(self,padding=(10,0,10,6));top.pack(fill='x')
        self.q=tk.StringVar();ttk.Label(top,text='Поиск:').pack(side='left');e=ttk.Entry(top,textvariable=self.q,width=36);e.pack(side='left',padx=5);e.bind('<Return>',lambda x:self.refresh());ttk.Button(top,text='Найти',command=self.refresh).pack(side='left')
        ttk.Button(top,text='Импорт фото по архивным №',command=lambda:self.app.screens['editor'].import_photos_by_number()).pack(side='left',padx=12)
        ttk.Label(top,text='Миниатюры:').pack(side='right');ttk.Scale(top,from_=48,to=110,variable=self.thumb_size,command=self.resize_thumbs).pack(side='right',padx=6)
        fr=ttk.Frame(self,padding=(10,0,10,8));fr.pack(fill='both',expand=True)
        cols=('no','desc','date','loc','author','source')
        self.tree=ttk.Treeview(fr,columns=cols,show='tree headings',selectmode='browse')
        self.tree.heading('#0',text='Фото');self.tree.column('#0',width=105,minwidth=90,stretch=False)
        specs=[('no','Архивный №',130),('desc','Описание',420),('date','Дата',100),('loc','Место',150),('author','Автор',150),('source','Источник',220)]
        for c,t,w in specs:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='w')
        ys=ttk.Scrollbar(fr,orient='vertical',command=self.tree.yview);xs=ttk.Scrollbar(fr,orient='horizontal',command=self.tree.xview);self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set)
        self.tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew');fr.rowconfigure(0,weight=1);fr.columnconfigure(0,weight=1)
        self.tree.bind('<Double-1>',self.open_clicked);self.tree.bind('<ButtonRelease-1>',self.click_open);self.tree.bind('<Motion>',self.hover);self.tree.bind('<Leave>',self.hide_hover)
        self.status()

    def refresh(self):
        self.rows=self.db.search(self.q.get(),'Все поля',True)
        for i in self.tree.get_children():self.tree.delete(i)
        self.thumb_refs={};size=max(48,int(self.thumb_size.get()));self.app.style.configure('Treeview',rowheight=max(58,size+8))
        for r in self.rows:
            tp=make_thumb(r['file_path'],r['id'],(size,int(size*.75)));img=''
            if tp and ImageTk:
                try:img=ImageTk.PhotoImage(Image.open(tp));self.thumb_refs[str(r['id'])]=img
                except Exception:img=''
            self.tree.insert('', 'end',iid=str(r['id']),text='',image=img,values=(r['archive_no'],safe_text(r['description']).replace('\n',' ')[:180],r['shot_date'],r['location'],r['author'],r['source']))
        self.status_var.set(f'Всего записей: {len(self.rows)}')
    def resize_thumbs(self,_=None):
        if hasattr(self,'_resize_job'):
            try:self.after_cancel(self._resize_job)
            except Exception:pass
        self._resize_job=self.after(250,self.refresh)
    def _row_path_at(self,y):
        iid=self.tree.identify_row(y)
        if not iid:return None,None
        r=self.db.get(int(iid));return r,r['file_path'] if r else None
    def click_open(self,e):
        region=self.tree.identify_region(e.x,e.y);col=self.tree.identify_column(e.x)
        if region=='tree' or col=='#0':
            r,p=self._row_path_at(e.y)
            if p:self.app.popup_photo(p,r['archive_no'])
    def open_clicked(self,e):
        r,p=self._row_path_at(e.y)
        if p:self.app.popup_photo(p,r['archive_no'])
    def hover(self,e):
        r,p=self._row_path_at(e.y)
        if not p or not os.path.exists(p) or Image is None:return self.hide_hover()
        if self.hover_after:
            try:self.after_cancel(self.hover_after)
            except Exception:pass
        self.hover_after=self.after(350,lambda:self.show_hover(p,e.x_root+18,e.y_root+18))
    def show_hover(self,p,x,y):
        self.hide_hover(cancel_only=True)
        try:
            img=Image.open(p);target_h=190;ratio=target_h/max(1,img.height);img=img.resize((max(1,int(img.width*ratio)),target_h))
            top=tk.Toplevel(self);top.overrideredirect(True);top.geometry(f'+{x}+{y}');ref=ImageTk.PhotoImage(img);lab=ttk.Label(top,image=ref,relief='solid');lab.image=ref;lab.pack();self.app.hover_popup=top
        except Exception:pass
    def hide_hover(self,_=None,cancel_only=False):
        if self.hover_after:
            try:self.after_cancel(self.hover_after)
            except Exception:pass
            self.hover_after=None
        if not cancel_only and self.app.hover_popup:
            try:self.app.hover_popup.destroy()
            except Exception:pass
            self.app.hover_popup=None


if __name__=='__main__':
    PhotoCatalogApp().mainloop()
